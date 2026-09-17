import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

def export_results_to_csv(results: list) -> bytes:
    output = io.StringIO()
    # Write BOM for Excel compatibility with UTF-8
    output.write('\ufeff')
    writer = csv.writer(output)
    writer.writerow([
        "ลำดับ (Row No)", 
        "วันที่เงินเข้า (Anchor Date)", 
        "ยอดเงินเข้า (Inflow Amount)", 
        "ยอดเงินออกรวมในกรอบเวลา (Outflow Total)", 
        "สัดส่วนเงินออก (Outflow Ratio)"
    ])
    
    for row in results:
        ratio_pct = f"{(row['outflow_ratio'] * 100):.2f}%"
        writer.writerow([
            row['row_no'], 
            row['anchor_date'], 
            row['inflow_amount'], 
            row['outflow_total'], 
            ratio_pct
        ])
        
    return output.getvalue().encode('utf-8')

def export_results_to_excel(results: list) -> bytes:
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "FAST_IN_OUT_3D Results"
    
    headers = [
        "ลำดับ (Row No)", 
        "วันที่เงินเข้า (Anchor Date)", 
        "ยอดเงินเข้า (Inflow Amount)", 
        "ยอดเงินออกรวมในกรอบเวลา (Outflow Total)", 
        "สัดส่วนเงินออก (Outflow Ratio)"
    ]
    ws.append(headers)
    for col in range(1, 6):
        ws.cell(row=1, column=col).font = Font(bold=True)
        ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
    
    for row in results:
        ratio_pct = f"{(row['outflow_ratio'] * 100):.2f}%"
        ws.append([
            row['row_no'], 
            row['anchor_date'], 
            row['inflow_amount'], 
            row['outflow_total'], 
            ratio_pct
        ])
        
    wb.save(output)
    return output.getvalue()

def export_results_to_pdf(results: list) -> bytes:
    output = io.BytesIO()
    # Use portrait or landscape depending on columns. Here columns are few, A4 portrait is fine.
    doc = SimpleDocTemplate(output, pagesize=A4)
    elements = []
    
    styles = getSampleStyleSheet()
    title = Paragraph("FAST_IN_OUT_3D Results", styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    # PDF using English to avoid font encoding issues since we don't have a Thai font loaded
    data = [["Row No", "Anchor Date", "Inflow Amount", "Outflow Total", "Outflow Ratio"]]
    for row in results:
        ratio_pct = f"{(row['outflow_ratio'] * 100):.2f}%"
        data.append([
            str(row['row_no']), 
            str(row['anchor_date']), 
            f"{row['inflow_amount']:,.2f}", 
            f"{row['outflow_total']:,.2f}", 
            ratio_pct
        ])
        
    t = Table(data, colWidths=[50, 90, 110, 110, 90])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('ALIGN', (0, 1), (0, -1), 'CENTER'),
        ('ALIGN', (2, 1), (-1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    elements.append(t)
    doc.build(elements)
    return output.getvalue()
